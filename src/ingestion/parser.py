import os
os.environ["TESSDATA_LANG"] = "eng"
os.environ["TESSERACT_LANG"] = "eng"
import csv
import logging
import hashlib
import openpyxl
import tiktoken
from unstructured.partition.auto import partition
from src.models import Chunk


logger = logging.getLogger(__name__)

MAX_CHUNK_TOKENS = 700

SUPPORTED_EXTENSIONS = {
    ".pdf", ".docx", ".doc", ".xlsx", ".xls", ".csv", ".html", ".htm", ".txt", ".md"
}

def generate_chunk_id(source_path: str, offset: int) -> str:
    return hashlib.sha256(f"{source_path}:{offset}".encode()).hexdigest()[:16]


def parse_file(path: str) -> list[Chunk]:
    """
    Convert any supported file into a list of Chunk objects using layout-aware parsing.
    """
    ext = os.path.splitext(path)[1].lower()
    
    if ext not in SUPPORTED_EXTENSIONS:
        logger.warning(f"Unsupported extension {ext} for file {path}. Skipping.")
        return []
        
    source_type = ext.lstrip('.')
    
    if ext in (".xlsx", ".xls"):
        return _parse_excel(path, source_type)
    elif ext == ".csv":
        return _parse_csv(path, source_type)
    else:
        return _parse_unstructured(path, source_type)


def _parse_excel(path: str, source_type: str) -> list[Chunk]:
    chunks = []
    wb = openpyxl.load_workbook(path, data_only=True)
    offset = 0
    abs_path = os.path.abspath(path)
    
    for sheet_name in wb.sheetnames:
        sheet = wb[sheet_name]
        
        md_table = []
        for row in sheet.iter_rows(values_only=True):
            row_vals = [str(cell) if cell is not None else "" for cell in row]
            md_table.append("| " + " | ".join(row_vals) + " |")
        
        if md_table:
            if len(md_table) > 1:
                header_cols = md_table[0].count("|") - 1
                separator = "| " + " | ".join(["---"] * header_cols) + " |"
                md_table.insert(1, separator)
                
            text = "\n".join(md_table)
            chunk = Chunk(
                chunk_id=generate_chunk_id(abs_path, offset),
                source_path=abs_path,
                source_type=source_type,
                page_number=None,
                section_header=sheet_name,
                text=text,
                element_type="Table",
                offset=offset
            )
            chunks.append(chunk)
            offset += 1
            
    return chunks


def _parse_csv(path: str, source_type: str) -> list[Chunk]:
    chunks = []
    offset = 0
    abs_path = os.path.abspath(path)
    
    with open(path, "r", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        fieldnames = reader.fieldnames or []
        
        if not fieldnames:
            return []
            
        header_row = "| " + " | ".join(fieldnames) + " |"
        separator = "| " + " | ".join(["---"] * len(fieldnames)) + " |"
        
        current_rows = []
        
        for row in reader:
            row_vals = [str(row.get(col, "")) for col in fieldnames]
            current_rows.append("| " + " | ".join(row_vals) + " |")
            
            if len(current_rows) == 100:
                text = "\n".join([header_row, separator] + current_rows)
                chunks.append(Chunk(
                    chunk_id=generate_chunk_id(abs_path, offset),
                    source_path=abs_path,
                    source_type=source_type,
                    page_number=None,
                    section_header=None,
                    text=text,
                    element_type="Table",
                    offset=offset
                ))
                offset += 1
                current_rows = []
                
        if current_rows:
            text = "\n".join([header_row, separator] + current_rows)
            chunks.append(Chunk(
                chunk_id=generate_chunk_id(abs_path, offset),
                source_path=abs_path,
                source_type=source_type,
                page_number=None,
                section_header=None,
                text=text,
                element_type="Table",
                offset=offset
            ))
            
    return chunks


def _parse_unstructured(path: str, source_type: str) -> list[Chunk]:
    # Use hi_res strategy for layout detection (e.g. PDFs)
    elements = partition(
        filename=path, 
        strategy="hi_res",
        languages=["eng"]
    )
    
    tokenizer = tiktoken.get_encoding("cl100k_base")
    chunks = []
    
    current_text_blocks = []
    current_token_count = 0
    current_section_header = None
    dominant_element_type = "NarrativeText"
    offset = 0
    current_page_number = None
    abs_path = os.path.abspath(path)
    
    def emit_chunk():
        nonlocal current_text_blocks, current_token_count, offset, dominant_element_type, current_page_number
        if current_text_blocks:
            text = "\n".join(current_text_blocks)
            chunks.append(Chunk(
                chunk_id=generate_chunk_id(abs_path, offset),
                source_path=abs_path,
                source_type=source_type,
                page_number=current_page_number,
                section_header=current_section_header,
                text=text,
                element_type=dominant_element_type,
                offset=offset
            ))
            offset += 1
            current_text_blocks = []
            current_token_count = 0
            dominant_element_type = "NarrativeText"
            current_page_number = None
            
    for el in elements:
        cat = el.category
        text = str(el)
        tokens = len(tokenizer.encode(text))
        
        page_num = None
        if hasattr(el, 'metadata') and el.metadata and hasattr(el.metadata, 'page_number'):
            page_num = el.metadata.page_number
        
        # 1. Title Boundary
        if cat == "Title":
            emit_chunk()
            current_section_header = text
            current_text_blocks.append(text)
            current_token_count += tokens
            current_page_number = page_num
            dominant_element_type = "Title"
            
        # 2. Table Boundary
        elif cat == "Table":
            emit_chunk()
            # Table is strictly its own chunk
            current_text_blocks.append(text)
            current_page_number = page_num
            dominant_element_type = "Table"
            emit_chunk()
            
        # 3. Token Count Boundary
        else:
            if current_token_count + tokens > MAX_CHUNK_TOKENS:
                emit_chunk()
                
            current_text_blocks.append(text)
            current_token_count += tokens
            if current_page_number is None:
                current_page_number = page_num
                
            if dominant_element_type not in ("Title", "Table"):
                dominant_element_type = cat
                
    emit_chunk()
    return chunks
